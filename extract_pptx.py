from pptx import Presentation

def extract_text_from_pptx(file_path):
    prs = Presentation(file_path)
    text = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        slide_text = f"Slide {slide_number}:\n"
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text += shape.text + "\n"
        text.append(slide_text)
    return "\n".join(text)

if __name__ == "__main__":
    file_path = "/home/student/projects/IoT/ВШЕ_IoT_Фомин (блок 2) v.2.0.pptx"
    extracted_text = extract_text_from_pptx(file_path)
    print(extracted_text)